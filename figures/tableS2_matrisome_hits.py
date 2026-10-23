#!/usr/bin/env python3
# Supplementary Table S2: the 23 matrisome proteins associated with matrix organization (Fig 6),
# grouped by direction, Division and Category, with a one-line function abridged from UniProtKB.
# In: results/proteomics/run_old_keyence/matrisome_intersection.csv, data/data_proteomics/uniprot/
# Out: results/thesis/supp_table_S2_matrisome_hits.xlsx (+ .csv and _legend.md)
import csv, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = "."
HITS = f"{ROOT}/results/proteomics/run_old_keyence/matrisome_intersection.csv"
UNI  = f"{ROOT}/data/data_proteomics/uniprot/matrisome_hits_uniprot_functions.json"
OUT  = f"{ROOT}/results/thesis/supp_table_S2_matrisome_hits"

# one-line functions, abridged from the UniProt "Function" text of each entry
FUNCTION = {
  # up in disorganized, core matrisome
  "ESM1":    "Secreted proteoglycan (endocan) that promotes angiogenic sprouting",
  "PXN":     "Paxillin: focal-adhesion adaptor linking integrins to the actin cytoskeleton (intracellular; see note)",
  "OIT3":    "Liver-enriched secreted glycoprotein of uncertain function; may act in hepatocyte development",
  "MATN3":   "Cartilage ECM glycoprotein that assembles extracellular filamentous networks",
  "CRIM1":   "Cysteine-rich glycoprotein that binds BMPs and modulates their processing and delivery; implicated in capillary formation",
  # up in disorganized, matrisome-associated: secreted factors
  "TNFSF4":  "Cytokine (OX40 ligand) that co-stimulates T-cell proliferation through TNFRSF4/OX40",
  "SCUBE3":  "BMP co-receptor and TGFBR2 ligand that promotes BMP and TGF-beta signaling",
  "ANGPTL4": "Angiopoietin-like factor that inhibits lipoprotein lipase and endothelial adhesion to the ECM, limiting angiogenesis",
  "EGFL7":   "Vascular secreted factor that regulates tubulogenesis and endothelial adhesion to the ECM",
  "INHBA":   "Activin A subunit; TGF-beta-family ligand that drives SMAD2/3 signaling, differentiation and tissue fibrosis",
  "PDGFC":   "PDGF-family growth factor; mitogen and chemoattractant for mesenchymal cells, drives fibroblast-to-myofibroblast transition",
  "VEGFC":   "VEGFR2/VEGFR3 ligand that drives lymphangiogenesis and angiogenesis",
  "SFRP2":   "Secreted frizzled-related protein that binds Wnts and modulates Wnt signaling",
  "WNT2B":   "Canonical Wnt/beta-catenin ligand of frizzled receptors",
  "WNT5A":   "Non-canonical Wnt ligand that stimulates cell migration and can suppress canonical Wnt signaling",
  "SCUBE2":  "Carrier for lipid-modified Sonic hedgehog enabling long-range SHH signaling; VEGFR2 co-receptor",
  "S100A16": "EF-hand calcium-binding protein that promotes adipocyte differentiation",
  # up in disorganized, matrisome-associated: ECM regulators
  "HPSE":    "Heparanase: endoglycosidase that cleaves heparan sulfate, degrading and remodeling the ECM and releasing bound growth factors",
  "ADAMTS4": "Aggrecanase-1: metalloproteinase that cleaves aggrecan and COMP in ECM turnover",
  # up in disorganized, matrisome-associated: ECM-affiliated
  "LGALS9":  "Galectin-9: beta-galactoside-binding lectin and TIM-3 ligand that induces Th1-cell death and regulates immune responses",
  "SEMA7A":  "GPI-anchored semaphorin that signals through integrins to promote focal adhesions, cell migration and monocyte cytokine production",
  # up in organized, core matrisome
  "SLIT3":   "Secreted axon-guidance glycoprotein; ROBO-receptor ligand acting as a cell-migration guidance cue",
  "SNED1":   "ECM glycoprotein that mediates RGD-dependent cell adhesion through integrins alpha5beta1 and alphaVbeta3",
}
# PXN matched the matrisome list as an alias of peroxidasin (PXDN), but P49023 is paxillin
NOTE = {"PXN": "Matched to the matrisome list through the symbol PXN, which the list carries as an alias of "
               "peroxidasin (PXDN); the quantified protein (P49023) is paxillin, an intracellular focal-adhesion "
               "protein rather than a secreted ECM glycoprotein."}

# row order follows the Results text: direction > Division > Category
ORDER = {("Up in disorganized", "Core matrisome"):       ["ESM1", "PXN", "OIT3", "MATN3", "CRIM1"],
         ("Up in disorganized", "Matrisome-associated"): ["TNFSF4", "SCUBE3", "ANGPTL4", "EGFL7", "INHBA", "PDGFC", "VEGFC",
                                                          "SFRP2", "WNT2B", "WNT5A", "SCUBE2", "S100A16", "HPSE", "ADAMTS4",
                                                          "LGALS9", "SEMA7A"],
         ("Up in organized",    "Core matrisome"):       ["SLIT3", "SNED1"]}
SECTION_TITLE = {"Up in disorganized": "Higher in disorganized matrix", "Up in organized": "Higher in organized matrix"}

# STEP 1. Join the hit table with the UniProt records
hits = {r["Genes"]: r for r in csv.DictReader(open(HITS))}
uni  = json.load(open(UNI))
recs = {r["gene_csv"]: r for r in uni["records"]}
assert len(hits) == 23 and set(hits) == set(recs) == set(FUNCTION)
assert sorted(sum(ORDER.values(), [])) == sorted(hits)
for g, r in hits.items():
    assert (r["direction"], r["Division"]) in ORDER and g in ORDER[(r["direction"], r["Division"])], g
    assert recs[g]["acc"] == r["UNIPROT"]

# STEP 2. Table rows
HEAD = ["Gene", "Protein (UniProt name)", "Matrisome category", "Function (abridged from UniProt)",
        "log2 FC (organized − disorganized)", "Adjusted p", "UniProt", "PubMed evidence cited by UniProt",
        "UniProt function text (verbatim)", "Note"]
rows = []   # (kind, value, direction)
for (direction, division), genes in ORDER.items():
    if not rows or rows[-1][0] != "section" and direction != rows[-1][2]:
        pass
    if not any(k == "section" and v == direction for k, v, *_ in rows):
        rows.append(("section", direction, direction))
    rows.append(("division", f"{division} ({len(genes)})", direction))
    for g in genes:
        r, u = hits[g], recs[g]
        # logFC_range is negated so positive = higher in organized matrix (as in Fig 6)
        rows.append(("protein", [g, u["name"], r["Category"], FUNCTION[g],
                                 round(-float(r["logFC_range"]), 2), float(r["adj.P.Val"]), u["acc"],
                                 "; ".join(u["pmids"]) if u["pmids"] else "(none cited in the entry)",
                                 u["function"], NOTE.get(g, "")], direction))

# table legend, in the style of the Table S1 legend
LEGEND = ("Supplementary Table S2. The 23 matrisome-annotated proteins associated with matrix organization "
          "(FDR < 0.05, Fig. 6A), grouped by direction and matrisome division (Naba et al., 2012). Functions are "
          f"abridged from UniProtKB (release {uni['uniprot_release']}; The UniProt Consortium, 2025) with the PubMed "
          "references UniProt cites; log2 FC is the change across the observed span of fiber spread (positive = higher "
          "in organized matrix), p is Benjamini–Hochberg adjusted.")
open(OUT + "_legend.md", "w").write("# Supplementary Table S2 legend\n\n" + LEGEND + "\n")

# STEP 3. Write the Excel and CSV tables
wb = Workbook(); ws = wb.active; ws.title = "Table S2"
thin = Side(style="thin", color="999999"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL = {"section": PatternFill("solid", fgColor="D9D9D9"),
        "Core matrisome": PatternFill("solid", fgColor="B4A7D6"), "Core matrisome_row": PatternFill("solid", fgColor="E4DFF2"),
        "Matrisome-associated": PatternFill("solid", fgColor="FFE599"), "Matrisome-associated_row": PatternFill("solid", fgColor="FFF2CC")}
ws.append(HEAD)
for c in ws[1]:
    c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="F2F2F2"); c.border = border
    c.alignment = Alignment(wrap_text=True, vertical="top")
current_division = None
for kind, val, direction in rows:
    if kind == "section":
        ws.append([SECTION_TITLE[val]]); r = ws.max_row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEAD))
        ws.cell(r, 1).font = Font(bold=True, size=12); ws.cell(r, 1).fill = FILL["section"]
    elif kind == "division":
        current_division = val.split(" (")[0]
        ws.append([val]); r = ws.max_row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEAD))
        ws.cell(r, 1).font = Font(bold=True); ws.cell(r, 1).fill = FILL[current_division]
    else:
        ws.append(val); r = ws.max_row
        ws.cell(r, 1).fill = FILL[current_division + "_row"]; ws.cell(r, 1).font = Font(bold=True)
        ws.cell(r, 6).number_format = "0.0E+00"
        for c in ws[r]:
            c.border = border; c.alignment = Alignment(wrap_text=True, vertical="top")
for i, w in enumerate([10, 34, 22, 60, 14, 11, 10, 34, 70, 40], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"

src = wb.create_sheet("Sources")
for line in [
    ["Legend", LEGEND],
    [""],
    ["Function", f"Abridged from the UniProtKB curated 'Function' annotation, release {uni['uniprot_release']} (retrieved {uni['fetched']}); the verbatim text and the PubMed IDs UniProt cites are given per row."],
    ["Cite as", "The UniProt Consortium. UniProt: the Universal Protein Knowledgebase in 2025. Nucleic Acids Res. 2025;53(D1):D609–D617. doi:10.1093/nar/gkae1010"],
    ["Matrisome division / category", "Naba A, et al. The matrisome: in silico definition and in vivo characterization by proteomics of normal and tumor extracellular matrices. Mol Cell Proteomics. 2012;11(4):M111.014647."],
    ["log2 FC", "Change in log2 intensity across the observed span of fiber spread (10.7–51.0° SD around the dominant direction), signed so that positive = higher in organized matrix; from results/proteomics/run_old_keyence/matrisome_intersection.csv (logFC_range, negated)."],
    ["Adjusted p", "Benjamini–Hochberg over the 7,578 tested proteins."],
    ["Note on PXN", NOTE["PXN"]],
]:
    src.append(line)
src.column_dimensions["A"].width = 30; src.column_dimensions["B"].width = 140
for row in src.iter_rows():
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(OUT + ".xlsx")

with open(OUT + ".csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["Group", "Division"] + HEAD)
    for kind, val, direction in rows:
        if kind == "division": current_division = val.split(" (")[0]
        if kind == "protein": w.writerow([SECTION_TITLE[direction], current_division] + val)
print("wrote", OUT + ".xlsx / .csv")
for kind, val, direction in rows:
    if kind != "protein": print(("  " if kind == "division" else "") + (val if kind == "division" else SECTION_TITLE[val]))
    else: print(f"    {val[0]:8s} {val[2]:24s} {val[4]:+.2f}  p={val[5]:.1e}  {val[3][:80]}")
